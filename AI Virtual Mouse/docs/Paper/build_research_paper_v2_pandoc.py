from __future__ import annotations

import re
import shutil
import subprocess
import zipfile
from html import escape
from pathlib import Path
from tempfile import TemporaryDirectory
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as docx_qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
PAPER_DIR = ROOT / "docs" / "Paper"
SOURCE_MD = PAPER_DIR / "research_paper_ai_virtual_mouse.md"
BIB = PAPER_DIR / "zotero_ai_virtual_mouse_export_all.bib"
V2_MD = PAPER_DIR / "research_paper_ai_virtual_mouse_v2.md"
CSL = PAPER_DIR / "elsevier_numeric_v2.csl"
REFERENCE_DOCX = PAPER_DIR / "research_paper_ai_virtual_mouse_v2_reference.docx"
PANDOC_RAW_DOCX = PAPER_DIR / "_research_paper_ai_virtual_mouse_v2_pandoc_raw.docx"
OUTPUT_DOCX = PAPER_DIR / "research_paper_ai_virtual_mouse_v2.docx"
ARTICLE_PLACEHOLDER = "__ARTICLE_INFO_ABSTRACT_TABLE__"


TITLE = (
    "Comparison of Baseline and Improved AI Virtual Mouse Hand-Control "
    "Pipelines Using a Point-and-Click Benchmark"
)
AUTHOR_LINE = "Raditya Rafif Pratama Sasmita"
AFFILIATION = "Politeknik Negeri Semarang, Semarang, Indonesia"


def extract_section(md: str, heading: str) -> str:
    pattern = rf"^## {re.escape(heading)}\s*\n(?P<body>.*?)(?=^## |\Z)"
    match = re.search(pattern, md, flags=re.M | re.S)
    if not match:
        raise RuntimeError(f"Section not found: {heading}")
    return match.group("body").strip()


def remove_front_matter_and_references(md: str) -> str:
    md = re.sub(r"^# .+?\n+", "", md, count=1, flags=re.S)
    md = re.sub(r"^## Abstract\s*\n.*?(?=^## 1\. Introduction)", "", md, flags=re.M | re.S)
    md = re.sub(r"^## References\s*\n.*\Z", "", md, flags=re.M | re.S).strip()
    return md


def split_keywords(keyword_line: str) -> list[str]:
    keyword_line = keyword_line.replace("**Keywords:**", "")
    keyword_line = keyword_line.replace("**Keywords**:", "")
    keyword_line = keyword_line.strip().strip(".")
    return [item.strip() for item in keyword_line.split(";") if item.strip()]


def paper_front_matter() -> tuple[str, list[str]]:
    md = SOURCE_MD.read_text(encoding="utf-8")
    abstract_block = extract_section(md, "Abstract")
    lines = [line.strip() for line in abstract_block.splitlines() if line.strip()]
    keyword_line = next((line for line in lines if line.startswith("**Keywords")), "")
    abstract = " ".join(line for line in lines if not line.startswith("**Keywords")).strip()
    keywords = split_keywords(keyword_line)
    return abstract, keywords


def write_v2_markdown() -> None:
    md = SOURCE_MD.read_text(encoding="utf-8")
    abstract, keywords = paper_front_matter()
    body = remove_front_matter_and_references(md)
    body = body.replace(
        "The structure of this paper follows the experimental style of Pan et al. "
        "[@pan_yolo-ecn_2026]:",
        "The structure of this paper follows the experimental style of the "
        "reference study [@pan_yolo-ecn_2026]:",
    )

    v2 = f"""---
bibliography: "zotero_ai_virtual_mouse_export_all.bib"
csl: "elsevier_numeric_v2.csl"
link-citations: true
reference-section-title: "References"
---

# {TITLE}

{AUTHOR_LINE}*

{AFFILIATION}

*Corresponding contributor: Raditya Rafif Pratama Sasmita.

{ARTICLE_PLACEHOLDER}

{body}
"""
    V2_MD.write_text(v2, encoding="utf-8", newline="\n")


def write_numeric_csl() -> None:
    CSL.write_text(
        """<?xml version="1.0" encoding="utf-8"?>
<style xmlns="http://purl.org/net/xbiblio/csl" version="1.0" class="in-text" default-locale="en-US">
  <info>
    <title>Elsevier-like Numeric for AI Virtual Mouse v2</title>
    <id>https://local/elsevier_numeric_ai_virtual_mouse_v2</id>
    <link href="https://local/elsevier_numeric_ai_virtual_mouse_v2" rel="self"/>
    <author><name>Codex</name></author>
    <updated>2026-06-03T00:00:00+00:00</updated>
  </info>
  <macro name="author">
    <names variable="author">
      <name initialize-with=". " delimiter=", " and="text"/>
      <label form="short" prefix=", "/>
      <substitute>
        <names variable="editor"/>
        <text variable="title"/>
      </substitute>
    </names>
  </macro>
  <macro name="title">
    <choose>
      <if type="article-journal paper-conference" match="any">
        <text variable="title" quotes="true"/>
      </if>
      <else>
        <text variable="title" font-style="italic"/>
      </else>
    </choose>
  </macro>
  <macro name="container">
    <choose>
      <if type="article-journal">
        <text variable="container-title" font-style="italic"/>
      </if>
      <else-if type="paper-conference">
        <text variable="container-title" prefix="in " font-style="italic"/>
      </else-if>
      <else>
        <text variable="container-title" font-style="italic"/>
      </else>
    </choose>
  </macro>
  <macro name="issued">
    <date variable="issued"><date-part name="year"/></date>
  </macro>
  <macro name="locators">
    <group delimiter=", ">
      <group delimiter=" ">
        <label variable="volume" form="short"/>
        <text variable="volume"/>
      </group>
      <group delimiter=" ">
        <label variable="issue" form="short"/>
        <text variable="issue"/>
      </group>
      <group delimiter=" ">
        <label variable="page" form="short"/>
        <text variable="page"/>
      </group>
    </group>
  </macro>
  <citation collapse="citation-number">
    <sort><key variable="citation-number"/></sort>
    <layout prefix="[" suffix="]" delimiter=", ">
      <text variable="citation-number"/>
    </layout>
  </citation>
  <bibliography second-field-align="flush" entry-spacing="0">
    <layout suffix=".">
      <text variable="citation-number" prefix="[" suffix="] "/>
      <group delimiter=", ">
        <text macro="author"/>
        <text macro="title"/>
        <text macro="container"/>
        <text macro="locators"/>
        <text macro="issued"/>
        <text variable="DOI" prefix="doi: "/>
        <text variable="URL"/>
      </group>
    </layout>
  </bibliography>
</style>
""",
        encoding="utf-8",
        newline="\n",
    )


def make_reference_docx() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.02)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(9)
    normal.paragraph_format.line_spacing = 1
    normal.paragraph_format.space_after = Pt(3)

    for style_name, size, bold in [
        ("Title", 18, True),
        ("Heading 1", 10, True),
        ("Heading 2", 9, True),
        ("Heading 3", 9, True),
        ("Body Text", 9, False),
        ("Caption", 8, False),
        ("Bibliography", 8, False),
    ]:
        try:
            style = styles[style_name]
        except KeyError:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold

    doc.add_paragraph("Reference document for Pandoc styles.")
    doc.save(REFERENCE_DOCX)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(docx_qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(docx_qn("w:fill"), fill)


def set_paragraph_bottom_border(paragraph, color: str = "000000", size: str = "8") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(docx_qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = p_bdr.find(docx_qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        p_bdr.append(bottom)
    bottom.set(docx_qn("w:val"), "single")
    bottom.set(docx_qn("w:sz"), size)
    bottom.set(docx_qn("w:space"), "1")
    bottom.set(docx_qn("w:color"), color)


def clear_paragraph_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(docx_qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def set_table_borders(table, color: str = "000000", size: str = "4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(docx_qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(docx_qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(docx_qn("w:val"), "single")
        element.set(docx_qn("w:sz"), size)
        element.set(docx_qn("w:space"), "0")
        element.set(docx_qn("w:color"), color)


def remove_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(docx_qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(docx_qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(docx_qn("w:val"), "nil")


def move_body_element(doc: Document, element, index: int) -> None:
    body = doc._body._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)
    body.insert(index, element)


def add_journal_style_header(doc: Document) -> None:
    top = doc.add_paragraph()
    top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = top.add_run("AI Virtual Mouse Research Article (2026)")
    run.font.name = "Times New Roman"
    run.font.size = Pt(6.5)
    run.font.color.rgb = RGBColor(0, 112, 192)

    header = doc.add_table(rows=1, cols=3)
    header.allow_autofit = False
    header.cell(0, 0).text = "ELSEVIER"
    header.cell(0, 1).text = "Computer Vision and Human-Computer Interaction\njournal-style research layout"
    header.cell(0, 2).text = "Prototype\nStudy"
    widths = [Inches(1.25), Inches(4.6), Inches(1.15)]
    for idx, width in enumerate(widths):
        header.cell(0, idx).width = width
        for paragraph in header.cell(0, idx).paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx != 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(7.5 if idx != 1 else 11)
                if idx == 0:
                    run.font.color.rgb = RGBColor(230, 126, 34)
                    run.bold = True
        if idx == 1:
            set_cell_shading(header.cell(0, idx), "EDEDED")
    remove_table_borders(header)

    thick_rule = doc.add_paragraph()
    set_paragraph_bottom_border(thick_rule, size="12")

    created = [top._p, header._tbl, thick_rule._p]
    for idx, element in enumerate(created):
        move_body_element(doc, element, idx)


def word_paragraph_text(element) -> str:
    return "".join(node.text or "" for node in element.iter(docx_qn("w:t")))


def replace_article_placeholder_with_table(doc: Document) -> None:
    abstract, keywords = paper_front_matter()
    body = doc._body._element
    placeholder = None
    placeholder_index = None
    accepted_placeholders = {ARTICLE_PLACEHOLDER, ARTICLE_PLACEHOLDER.strip("_")}
    for index, child in enumerate(list(body)):
        if child.tag == docx_qn("w:p") and word_paragraph_text(child).strip() in accepted_placeholders:
            placeholder = child
            placeholder_index = index
            break
    if placeholder is None or placeholder_index is None:
        raise RuntimeError("Article info placeholder not found in Pandoc DOCX")

    table = doc.add_table(rows=1, cols=2)
    table.allow_autofit = False
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.width = Inches(2.05)
    right.width = Inches(4.85)

    left.paragraphs[0].text = "A R T I C L E   I N F O"
    left.paragraphs[0].paragraph_format.space_after = Pt(10)
    left.paragraphs[0].runs[0].bold = False
    left.paragraphs[0].runs[0].font.size = Pt(7.5)
    left.paragraphs[0].runs[0].font.name = "Times New Roman"

    kw = left.add_paragraph("Keywords:")
    kw.runs[0].italic = True
    kw.runs[0].font.size = Pt(7)
    kw.runs[0].font.name = "Times New Roman"
    kw.paragraph_format.space_after = Pt(2)
    for keyword in keywords:
        paragraph = left.add_paragraph(keyword)
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.size = Pt(6.5)
            run.font.name = "Times New Roman"

    right.paragraphs[0].text = "A B S T R A C T"
    right.paragraphs[0].paragraph_format.space_after = Pt(8)
    right.paragraphs[0].runs[0].font.size = Pt(7.5)
    right.paragraphs[0].runs[0].font.name = "Times New Roman"
    abstract_p = right.add_paragraph(abstract)
    abstract_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    abstract_p.paragraph_format.line_spacing = 1
    abstract_p.paragraph_format.space_after = Pt(0)
    for run in abstract_p.runs:
        run.font.size = Pt(7)
        run.font.name = "Times New Roman"

    set_table_borders(table, size="4")

    body.remove(placeholder)
    move_body_element(doc, table._tbl, placeholder_index)


def run_pandoc() -> None:
    cmd = [
        "pandoc",
        "--from=markdown+raw_html",
        "--citeproc",
        f"--bibliography={BIB.name}",
        f"--csl={CSL.name}",
        f"--reference-doc={REFERENCE_DOCX}",
        "--metadata",
        "link-citations=true",
        str(V2_MD.name),
        "--output",
        str(PANDOC_RAW_DOCX.name),
    ]
    subprocess.run(cmd, cwd=PAPER_DIR, check=True)


def style_docx() -> None:
    doc = Document(PANDOC_RAW_DOCX)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.02)
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    add_journal_style_header(doc)
    replace_article_placeholder_with_table(doc)

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        paragraph_format = paragraph.paragraph_format
        paragraph_format.line_spacing = 1
        if text == TITLE:
            clear_paragraph_border(paragraph)
            paragraph.style = doc.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph_format.space_before = Pt(28)
            paragraph_format.space_after = Pt(10)
            for run in paragraph.runs:
                run.bold = False
                run.underline = False
                run.font.name = "Times New Roman"
                run.font.size = Pt(13)
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif text in {AUTHOR_LINE + "*", AUTHOR_LINE, AFFILIATION}:
            paragraph.style = doc.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph_format.space_after = Pt(2)
            paragraph_format.space_before = Pt(0)
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(0, 0, 0)
                run.bold = False
        elif text.startswith("*Corresponding contributor"):
            paragraph.style = doc.styles["Normal"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph_format.space_before = Pt(0)
            paragraph_format.space_after = Pt(8)
            set_paragraph_bottom_border(paragraph, size="6")
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(7)
                run.font.italic = True
                run.font.color.rgb = RGBColor(0, 0, 0)
        elif text.startswith("Table ") or text.startswith("Fig."):
            paragraph.style = doc.styles["Caption"]
            paragraph_format.space_before = Pt(4)
            paragraph_format.space_after = Pt(2)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            if run.font.color.rgb is None:
                run.font.color.rgb = RGBColor(0, 0, 0)

    for table_index, table in enumerate(doc.tables):
        if table_index == 0:
            continue
        table.allow_autofit = True
        set_table_borders(table, size="4")
        table_font_size = Pt(7 if table_index == 1 else 8)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.line_spacing = 1
                    paragraph.paragraph_format.space_after = Pt(2)
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = table_font_size
                        run.font.color.rgb = RGBColor(0, 0, 0)

    doc.save(PANDOC_RAW_DOCX)


def qn(tag: str) -> str:
    prefix, name = tag.split(":")
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    }[prefix]
    return f"{{{ns}}}{name}"


def sect_pr(cols: int) -> ET.Element:
    sect = ET.Element(qn("w:sectPr"))
    ET.SubElement(sect, qn("w:type"), {qn("w:val"): "continuous"})
    ET.SubElement(sect, qn("w:pgSz"), {qn("w:w"): "11906", qn("w:h"): "15874"})
    ET.SubElement(
        sect,
        qn("w:pgMar"),
        {
            qn("w:top"): "792",
            qn("w:right"): "792",
            qn("w:bottom"): "792",
            qn("w:left"): "792",
            qn("w:header"): "360",
            qn("w:footer"): "360",
            qn("w:gutter"): "0",
        },
    )
    ET.SubElement(sect, qn("w:cols"), {qn("w:space"): "360", qn("w:num"): str(cols)})
    return sect


def paragraph_text(paragraph: ET.Element) -> str:
    return "".join(node.text or "" for node in paragraph.iter(qn("w:t")))


def patch_columns() -> None:
    ET.register_namespace("w", "http://schemas.openxmlformats.org/wordprocessingml/2006/main")
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")

    with TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(PANDOC_RAW_DOCX, "r") as zin:
            zin.extractall(tmp_path)

        doc_xml = tmp_path / "word" / "document.xml"
        tree = ET.parse(doc_xml)
        root = tree.getroot()
        body = root.find(qn("w:body"))
        if body is None:
            raise RuntimeError("DOCX body not found")

        children = list(body)
        intro_index = None
        for idx, child in enumerate(children):
            if child.tag == qn("w:p") and paragraph_text(child).strip().startswith("1. Introduction"):
                intro_index = idx
                break
        if intro_index is None:
            raise RuntimeError("Could not locate Introduction paragraph for section break")

        section_break = ET.Element(qn("w:p"))
        ppr = ET.SubElement(section_break, qn("w:pPr"))
        ppr.append(sect_pr(1))
        body.insert(intro_index, section_break)

        final_sect = body.find(qn("w:sectPr"))
        if final_sect is None:
            final_sect = sect_pr(2)
            body.append(final_sect)
        else:
            for node in list(final_sect):
                if node.tag in {qn("w:cols"), qn("w:pgSz"), qn("w:pgMar")}:
                    final_sect.remove(node)
            replacement = sect_pr(2)
            for node in list(replacement):
                if node.tag != qn("w:type"):
                    final_sect.append(node)

        tree.write(doc_xml, encoding="utf-8", xml_declaration=True)

        with zipfile.ZipFile(OUTPUT_DOCX, "w", zipfile.ZIP_DEFLATED) as zout:
            for path in tmp_path.rglob("*"):
                if path.is_file():
                    zout.write(path, path.relative_to(tmp_path).as_posix())


def main() -> None:
    write_v2_markdown()
    write_numeric_csl()
    make_reference_docx()
    run_pandoc()
    style_docx()
    patch_columns()
    print(f"Wrote {V2_MD}")
    print(f"Wrote {REFERENCE_DOCX}")
    print(f"Wrote {OUTPUT_DOCX}")


if __name__ == "__main__":
    main()
